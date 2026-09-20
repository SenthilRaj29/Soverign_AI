import os
from typing import Dict, Any, List

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
except ImportError:
    SimpleDocTemplate = None

class ReportGenerator:
    def __init__(self, output_dir: str = "e:/Soverign_AI/reports/generated"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    def generate_docx_report(self, task_id: str, title: str, summary: str, findings: List[str], recommendations: List[str], sources: List[str]) -> str:
        filepath = os.path.join(self.output_dir, f"{task_id}_maintenance_report.docx")
        if Document is None:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(f"# {title}\n\n## Summary\n{summary}\n\n## Findings\n" + "\n".join(findings))
            return filepath
            
        doc = Document()
        doc.add_heading(title, 0)
        doc.add_heading("Executive Summary", level=1)
        doc.add_paragraph(summary)
        doc.add_heading("Key Findings & Sensor Analysis", level=1)
        for finding in findings:
            doc.add_paragraph(f"• {finding}")
        doc.add_heading("Safety & Maintenance Recommendations", level=1)
        for rec in recommendations:
            doc.add_paragraph(f"1. {rec}")
        doc.add_heading("Verified Sources & Audit Reference", level=1)
        for src in sources:
            doc.add_paragraph(f"Source Reference: {src}")
        doc.save(filepath)
        return filepath

    def generate_pdf_report(self, task_id: str, title: str, summary: str, findings: List[str], recommendations: List[str], sources: List[str]) -> str:
        filepath = os.path.join(self.output_dir, f"{task_id}_maintenance_report.pdf")
        if SimpleDocTemplate is None:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(f"# {title}\n\n## Summary\n{summary}\n\n## Findings\n" + "\n".join(findings))
            return filepath

        doc = SimpleDocTemplate(filepath, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#1E293B'),
            spaceAfter=14
        )
        story.append(Paragraph(title, title_style))
        story.append(Spacer(1, 10))

        story.append(Paragraph("<b>EXECUTIVE SUMMARY</b>", styles['Heading2']))
        story.append(Paragraph(summary, styles['BodyText']))
        story.append(Spacer(1, 10))

        story.append(Paragraph("<b>KEY FINDINGS & SENSOR ANALYSIS</b>", styles['Heading2']))
        for f in findings:
            story.append(Paragraph(f"• {f}", styles['BodyText']))
        story.append(Spacer(1, 10))

        story.append(Paragraph("<b>RECOMMENDED ACTIONS</b>", styles['Heading2']))
        for r in recommendations:
            story.append(Paragraph(f"<b>ACTION:</b> {r}", styles['BodyText']))
        story.append(Spacer(1, 10))

        story.append(Paragraph("<b>AUDIT & KNOWLEDGE SOURCES</b>", styles['Heading2']))
        for s in sources:
            story.append(Paragraph(f"Reference: {s}", styles['Italic']))

        doc.build(story)
        return filepath
